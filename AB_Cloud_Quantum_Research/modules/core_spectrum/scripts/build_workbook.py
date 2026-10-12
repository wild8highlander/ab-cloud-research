#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""build_workbook.py — сборка tables/rp_workbook.xlsx (openpyxl).

Листы: Headline, Programs V1–V19, Theorems, Scorecard, Data files.
Запуск: python3 scripts/build_workbook.py
"""
import csv
import glob
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, 'data')

HEADER_FILL = PatternFill('solid', fgColor='3B4F58')
STRIPE = PatternFill('solid', fgColor='EBEDEE')
THIN = Side(style='thin', color='C7D3D9')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def sheet(wb, name, header, rows, widths=None, note=None):
    ws = wb.create_sheet(name)
    ws.append(header)
    for c in ws[1]:
        c.fill = HEADER_FILL
        c.font = Font(color='FFFFFF', bold=True, size=9)
        c.alignment = Alignment(vertical='center')
        c.border = BORDER
    for i, row in enumerate(rows):
        ws.append([str(x) for x in row])
        for c in ws[i + 2]:
            c.font = Font(size=9)
            c.border = BORDER
            if i % 2 == 1:
                c.fill = STRIPE
    for j, w in enumerate(widths or [22] * len(header)):
        ws.column_dimensions[chr(65 + j)].width = w
    ws.freeze_panes = 'A2'
    if note:
        r = ws.max_row + 2
        ws.cell(row=r, column=1, value=note).font = Font(italic=True,
                                                         size=8)
    return ws


def main():
    wb = Workbook()
    wb.remove(wb.active)

    # Headline
    head = [
        ('1', 'Th1 эрмитовость', '‖H−H†‖ = 0.0 (точно)', 'V4'),
        ('2', 'Th2 самосопряжённость', 'дрейф унитарности 2.2e-16', 'V8'),
        ('3', 'Th3 сертификаты', '5/5 интервалов содержат λ', 'V12'),
        ('4', 'Th4 спектральная формула Ли', '|Δ| двух путей = 4.4e-11', 'V15'),
        ('5', 'λ_n > 0', 'n ≤ 3000, хвост ≤ 0.115', 'V15'),
        ('6', 'Th5 блочная редукция', 'union err = 7.1e-15', 'V9'),
        ('7', 'Th6 следовая формула', '|ΔK| = 6.9e-6', 'V7'),
        ('8', 'PC0 позитивный контроль', 'r_EF = +0.9996 (resid 0.027)', 'V7'),
        ('9', 'разрыв затвор↔гребёнка', '≈ 0.63 порядка по невязке', 'V7'),
        ('10', 'слепой контроль', 'r_test = 0.006, z = 0.09', 'V1'),
        ('11', 'GUE-класс AB', 'D_KS = 0.050 (эталон GUE 0.062)', 'V2'),
        ('12', 'скан σ', 'минимум в ½ отсутствует', 'V3'),
        ('13', 'предел L→∞', 'e(L) ≈ L^−0.86, R² = 0.80', 'V5'),
        ('14', 'спектральное тождество', 'затворы НЕТ; H_ξ resid 5.8e-10', 'V6'),
        ('15', 'обратная задача', 'восстановление 0.077, rank 6/12', 'V10'),
        ('16', 'универсальность', '⟨r⟩ выживает во всех режимах', 'V11'),
        ('17', 'формальный слой', 'RH: blocked, 7 блокеров', 'V13'),
        ('18', 'аудит', '19/19 claims по схеме', 'V14'),
        ('19', 'полный гамильтониан', 'тождество НЕ замыкается: z = 2.49', 'V16'),
        ('20', 'ПК Z_id ↔ log|ζ|', 'r = 1.0 (алгебраическая функция)', 'V16'),
        ('21', 'электрон: динамика', 'ΔE = 1.8e-15 (честная)', 'V17'),
        ('22', 'электрон: направление', 'ζ-облако ×1.62; r_phase −0.96', 'V17'),
        ('23', 'AB-процессор: ворота', '{X,Y,Z,H,S,T} точные (1−2e-16)', 'V18'),
        ('24', 'AB-процессор: игра', 'AB-Pong 0:1, реплей бит-в-бит', 'V18'),
        ('25', 'Сшивка ветвей (R2)', 'сертификат ALGEBRAIC_CERTIFICATE: фазы 1e-40, чередование 100%, когерентность 0.9995', 'V19'),
    ]
    sheet(wb, 'Headline', ['#', 'Результат', 'Значение', 'Кампания'], head,
          [5, 34, 42, 11],
          'AB Research Program 1.2.1 · все числа из закоммиченных '
          'вердиктов data/')

    # Programs V1–V19
    rows = []
    for k in range(1, 20):
        d = json.load(open(os.path.join(DATA, f'v{k:02d}_verdict.json')))
        m = d.get('meta', {})
        rows.append((f'V{k:02d}', m.get('campaign', ''), m.get('date'),
                     f"{m.get('runtime_s', '?')} s",
                     str(d.get('verdict_note', ''))[:160]))
    sheet(wb, 'Programs', ['ID', 'Кампания', 'Дата', 'Время', 'Вердикт'],
          rows, [7, 10, 12, 9, 90])

    # Theorems
    th = [
        ('Th1', 'V4', 'доказано', 'эрмитовость H_L, spec ⊂ ℝ', '‖H−H†‖ = 0'),
        ('Th2', 'V8', 'доказано (конечномерно)', 'дефекты (0,0)',
         'дрейф 2.2e-16'),
        ('Th3', 'V12', 'доказано + верифицировано', 'сертификаты λ',
         '5/5 valid'),
        ('Th4', 'V15', 'доказано + верифицировано', 'λ_n = Tr g_n(H_ξ)',
         '|Δ| 4.4e-11'),
        ('Th5', 'V9', 'доказано + верифицировано', 'блочная редукция',
         '7.1e-15'),
        ('Th6', 'V7', 'доказано + верифицировано', 'Гуинанд–Вейль',
         '|ΔK| 6.9e-6'),
    ]
    sheet(wb, 'Theorems', ['Теорема', 'Кампания', 'Статус', 'Содержание',
                           'Ключевое число'], th, [10, 10, 26, 30, 14])

    # Scorecard
    sc = [
        ('V1', 'Слепой спектр', 'измерено', 'r_test в нулевой зоне (z=0.09)'),
        ('V2', 'Null-модели', 'измерено', 'AB ∈ GUE; N3 вне класса'),
        ('V3', 'Скан σ', 'измерено', 'минимума в ½ нет'),
        ('V4', 'Паспорт Th1', 'доказано', 'эрмитовость точная'),
        ('V5', 'Предел', 'измерено', 'p ≈ 0.86, R² = 0.80'),
        ('V6', 'Тождество', 'измерено', 'затворы НЕТ; H_ξ ДА'),
        ('V7', 'След Th6', 'доказано', 'PC0 +0.9996; затворы — нет'),
        ('V8', 'Самосопряжённость Th2', 'доказано', '(0,0); унитарность'),
        ('V9', 'Редукция Th5', 'доказано', 'блоки = спектр (7e-15)'),
        ('V10', 'Обратная задача', 'измерено', 'rank-limited'),
        ('V11', 'Универсальность', 'измерено', 'GUE выживает'),
        ('V12', 'Сертификаты Th3', 'доказано', 'гарантированные интервалы'),
        ('V13', 'Формальный слой', 'измерено', 'RH: blocked'),
        ('V14', 'Аудит', 'измерено', '19/19 claims'),
        ('V15', 'Критерий Ли Th4', 'доказано', 'λ_n > 0, n ≤ 3000'),
        ('V16', 'Полный гамильтониан', 'измерено', 'тождество не замыкается (z=2.49); ПК r=1.0'),
        ('V17', 'Электрон на Re(s)=½', 'измерено + конъектура', 'ζ-облако направляет ×1.62; ΔE 1.8e-15'),
        ('V18', 'AB-процессор', 'измерено', 'ворота точные; C=0.984; AB-Pong 0:1'),
        ('V19', 'Сшивка ветвей', 'ALGEBRAIC_CERTIFICATE', 'фазовый закон 1e-40; арг. замыкание 138=138; когерентность 0.9995 (z=43.4)'),
    ]
    sheet(wb, 'Scorecard', ['ID', 'Направление', 'Статус', 'Итог'], sc,
          [7, 26, 14, 52])

    # Data files
    files = sorted(os.path.basename(f) for f in glob.glob(
        os.path.join(DATA, '*')))
    rows = [(f, os.path.getsize(os.path.join(DATA, f))) for f in files]
    sheet(wb, 'Data files', ['Файл', 'Байт'], rows, [40, 12])

    out = os.path.join(ROOT, 'tables', 'rp_workbook.xlsx')
    os.makedirs(os.path.dirname(out), exist_ok=True)
    wb.save(out)
    print(f'XLSX OK: {out}')


if __name__ == '__main__':
    main()
